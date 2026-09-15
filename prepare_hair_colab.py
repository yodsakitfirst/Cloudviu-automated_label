"""Build a runtime-local Colab Enterprise package for the HAIR dataset."""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Mapping, Sequence


_PRODUCT_EXTENSIONS = frozenset({".jpg", ".jpeg", ".png"})


@dataclass(frozen=True)
class HairSku:
    class_id: int
    barcode: str
    sku_name: str


def normalize_product_name(value: str) -> str:
    """Normalize a product name only for exact, deterministic matching."""
    return " ".join(unicodedata.normalize("NFC", value).split()).casefold()


def _barcode_text(value: object, row_number: int) -> str:
    if isinstance(value, bool):
        raise ValueError(f"Invalid barcode in Sheet2 row {row_number}")
    if isinstance(value, int):
        barcode = str(value)
    elif isinstance(value, float) and value.is_integer():
        barcode = str(int(value))
    else:
        barcode = str(value or "").strip()
    if not barcode.isascii() or not barcode.isdigit() or len(barcode) not in {13, 14}:
        raise ValueError(f"Invalid barcode in Sheet2 row {row_number}: {barcode!r}")
    return barcode


def load_sheet2_skus(path: Path, expected_count: int = 89) -> list[HairSku]:
    """Load the permanent class registry from Sheet2 rows 3 onward."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read the SKU workbook") from exc

    workbook_path = Path(path).resolve()
    if not workbook_path.is_file():
        raise ValueError(f"Workbook does not exist: {workbook_path}")
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Cannot read workbook {workbook_path}: {exc}") from exc
    try:
        if "Sheet2" not in workbook.sheetnames:
            raise ValueError("Workbook is missing Sheet2")
        sheet = workbook["Sheet2"]
        headers = tuple(str(sheet.cell(2, column).value or "").strip() for column in (1, 2))
        if headers != ("Barcode", "Product Name"):
            raise ValueError(
                "Sheet2 row 2 must start with 'Barcode' and 'Product Name'"
            )

        source_rows: list[tuple[int, object, object]] = []
        for row_number in range(3, sheet.max_row + 1):
            barcode_value = sheet.cell(row_number, 1).value
            name_value = sheet.cell(row_number, 2).value
            if barcode_value is None and name_value is None:
                continue
            source_rows.append((row_number, barcode_value, name_value))
    finally:
        workbook.close()

    if len(source_rows) != expected_count:
        raise ValueError(f"Expected {expected_count} SKUs, found {len(source_rows)}")

    skus: list[HairSku] = []
    barcodes: set[str] = set()
    names: set[str] = set()
    for class_id, (row_number, barcode_value, name_value) in enumerate(source_rows):
        barcode = _barcode_text(barcode_value, row_number)
        sku_name = unicodedata.normalize("NFC", str(name_value or "").strip())
        name_key = normalize_product_name(sku_name)
        if not sku_name:
            raise ValueError(f"Blank product name in Sheet2 row {row_number}")
        if barcode in barcodes:
            raise ValueError(f"Duplicate barcode in Sheet2 row {row_number}: {barcode}")
        if name_key in names:
            raise ValueError(f"Duplicate product name in Sheet2 row {row_number}: {sku_name}")
        barcodes.add(barcode)
        names.add(name_key)
        skus.append(HairSku(class_id, barcode, sku_name))
    return skus


def load_reference_overrides(path: Path) -> dict[str, str]:
    """Load explicit workbook-name to source-filename exceptions."""
    try:
        import yaml
    except ImportError as exc:
        raise RuntimeError("PyYAML is required to read reference overrides") from exc

    override_path = Path(path).resolve()
    try:
        payload = yaml.safe_load(override_path.read_text(encoding="utf-8"))
    except (OSError, yaml.YAMLError) as exc:
        raise ValueError(f"Cannot read reference overrides {override_path}: {exc}") from exc
    if not isinstance(payload, dict) or not isinstance(payload.get("overrides"), dict):
        raise ValueError("Reference overrides must contain an 'overrides' mapping")
    overrides = payload["overrides"]
    if any(
        not isinstance(product_name, str)
        or not product_name.strip()
        or not isinstance(filename, str)
        or not filename.strip()
        for product_name, filename in overrides.items()
    ):
        raise ValueError("Reference overrides must map string product names to string filenames")
    return {product_name.strip(): filename.strip() for product_name, filename in overrides.items()}


def match_product_references(
    skus: Sequence[HairSku],
    product_dir: Path,
    overrides: Mapping[str, str],
) -> dict[int, Path]:
    """Match every SKU to exactly one reference without fuzzy matching."""
    root = Path(product_dir).resolve()
    if not root.is_dir():
        raise ValueError(f"Product image directory does not exist: {root}")
    images = sorted(
        (
            path.resolve()
            for path in root.iterdir()
            if path.is_file() and path.suffix.casefold() in _PRODUCT_EXTENSIONS
        ),
        key=lambda path: (path.name.casefold(), path.name),
    )
    by_name: dict[str, Path] = {}
    for image in images:
        key = normalize_product_name(image.stem)
        if key in by_name:
            raise ValueError(
                f"Ambiguous product reference names: {by_name[key].name} and {image.name}"
            )
        by_name[key] = image

    normalized_overrides: dict[str, tuple[str, str]] = {}
    for product_name, filename in overrides.items():
        key = normalize_product_name(product_name)
        if key in normalized_overrides:
            raise ValueError(f"Duplicate normalized override for product: {product_name}")
        normalized_overrides[key] = (product_name, filename)

    matched: dict[int, Path] = {}
    used_override_keys: set[str] = set()
    used_images: set[Path] = set()
    for sku in skus:
        name_key = normalize_product_name(sku.sku_name)
        if name_key in normalized_overrides:
            _, filename = normalized_overrides[name_key]
            candidate = (root / filename).resolve()
            if candidate.parent != root or not candidate.is_file():
                raise ValueError(
                    f"Reference override for {sku.sku_name!r} does not resolve to a file: {filename}"
                )
            if candidate.suffix.casefold() not in _PRODUCT_EXTENSIONS:
                raise ValueError(f"Unsupported product reference extension: {candidate.name}")
            used_override_keys.add(name_key)
        else:
            candidate = by_name.get(name_key)
            if candidate is None:
                raise ValueError(f"Missing reference for product: {sku.sku_name}")
        if candidate in used_images:
            raise ValueError(f"Product reference is assigned more than once: {candidate.name}")
        used_images.add(candidate)
        matched[sku.class_id] = candidate

    unused_overrides = sorted(set(normalized_overrides) - used_override_keys)
    if unused_overrides:
        original_names = [normalized_overrides[key][0] for key in unused_overrides]
        raise ValueError(f"Unused override product names: {original_names}")
    unused_images = [image.name for image in images if image not in used_images]
    if unused_images:
        raise ValueError(f"Unmatched product reference images: {unused_images}")
    return matched


def ascii_reference_name(sku: HairSku, source: Path) -> str:
    """Return the deterministic English-safe reference filename."""
    suffix = Path(source).suffix.casefold()
    if suffix not in _PRODUCT_EXTENSIONS:
        raise ValueError(f"Unsupported product reference extension: {source}")
    filename = f"class_{sku.class_id:03d}_{sku.barcode}{suffix}"
    if not filename.isascii():
        raise ValueError(f"Prepared reference name is not ASCII: {filename}")
    return filename
